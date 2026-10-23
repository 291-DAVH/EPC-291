"""Resolve sibling plugins from the EPC plugin hub (platform/hub/<plugin-id>/).

- aea-calc         -> calculation formulas (scripts/voltage_drop.py, ...)
- iec-symbols      -> symbol library (symbols/sld-lv-symbol)
- iram-titleblock  -> title block and IRAM sheet (scripts/, templates/)

Hub location: EPC_HUB env var, otherwise the parent of this plugin folder (platform/hub).
Replaces the old vendor/*.plugin copies: there is exactly one copy of each plugin.
"""
import os, sys, tempfile, subprocess

HERE = os.path.dirname(os.path.abspath(__file__))
PLUGIN_ROOT = os.path.dirname(HERE)
HUB = os.environ.get('EPC_HUB') or os.path.dirname(PLUGIN_ROOT)
CACHE = os.path.join(tempfile.gettempdir(), 'epc_sld_work')

_HUB_ID = {'iec60617-symbol-generator': 'iec-symbols', 'elec-calc-generator': 'aea-calc',
           'iram-titleblock-plugin': 'iram-titleblock'}


def _dir(name):
    d = os.path.join(HUB, _HUB_ID.get(name, name))
    if not os.path.isdir(d):
        raise FileNotFoundError(f'Plugin not found in hub: {d}')
    return d


def iec_lib():
    """Carpeta con los símbolos SLD BT (_catalog.json + *.svg)."""
    return os.path.join(_dir('iec60617-symbol-generator'), 'symbols', 'sld-lv-symbol')


def add_elec_calc_path():
    """Agrega scripts de elec-calc-generator al path y devuelve la carpeta."""
    p = os.path.join(_dir('elec-calc-generator'), 'scripts')
    if p not in sys.path:
        sys.path.insert(0, p)
    return p


ROTULO_DIR = os.environ.get('SLD_ROTULO_DIR') or os.path.join(PLUGIN_ROOT, 'examples', 'titleblock')   # input-titleblocks.xlsx + logos
# Columnas de la plantilla que son datos de ejemplo del plugin: se vacían al crear el Excel del proyecto
_EJEMPLO = ('Empresa:', 'Cliente:', 'Proyectado por', 'Verificado por', 'Validado por', 'Descripción', 'Ubicación')


def _crear_excel(xlsx, datos):
    """Crea input-titleblocks.xlsx desde la plantilla del plugin (sin datos de ejemplo)."""
    import openpyxl
    src = os.path.join(_dir('iram-titleblock-plugin'), 'templates', 'input-titleblocks.xlsx')
    wb = openpyxl.load_workbook(src)
    ws = wb['Láminas']
    hdr = {c.value: c.column for c in ws[1]}
    for r in range(2, ws.max_row + 1):
        for h, col in hdr.items():
            if h and str(h).startswith(_EJEMPLO):
                ws.cell(r, col).value = None
    # datos que el SLD conoce, sólo para la fila A1 (se escriben una vez, al crear el archivo)
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, hdr['Formato (HOJA)']).value == 'A1':
            for h, v in datos.items():
                if h in hdr and v:
                    ws.cell(r, hdr[h]).value = v
    # 'Archivo de salida' es una fórmula sin valor en caché: el generador la lee con data_only, así que se escribe el valor
    fi, fa, fm = hdr['Formato (HOJA)'], hdr['Archivo de salida'], hdr['Lista de modificaciones']
    for r in range(2, ws.max_row + 1):
        f = ws.cell(r, fi).value
        if f:
            mod = ws.cell(r, fm).value == 'SI'
            ws.cell(r, fa).value = f"sheet-IRAM-{f}{'-mod' if mod else ''}.svg"
    os.makedirs(os.path.dirname(xlsx), exist_ok=True)
    wb.save(xlsx)


def _excel():
    xlsx = os.path.join(ROTULO_DIR, 'input-titleblocks.xlsx')
    if not os.path.isfile(xlsx):
        _crear_excel(xlsx, {})
    return xlsx


def sincronizar(nombres, comunes, fmt='A1'):
    """Deja input-titleblocks.xlsx con UNA fila por hoja (nombres = planos propios de cada hoja).

    Los datos comunes (empresa, cliente, personas...) se toman de la primera fila `fmt` del Excel.
    Cada fila recibe: hoja N°, plano propio, página n / m, denominación y proyecto de esa hoja.
    """
    import openpyxl
    xlsx = _excel()
    wb = openpyxl.load_workbook(xlsx)
    ws = wb['Láminas']
    hdr = {c.value: c.column for c in ws[1]}
    base = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, hdr['Formato (HOJA)']).value == fmt)
    vals = [c.value for c in ws[base]]
    if ws.max_row > 1: ws.delete_rows(2, ws.max_row - 1)
    n = len(nombres)
    for i, nombre in enumerate(nombres, 1):
        for c, v in enumerate(vals, 1): ws.cell(i + 1, c).value = v
        datos = dict(comunes)
        if str(nombre).endswith('_TABLA'):   # hoja de tabla de cálculo
            datos['Denominación (línea 2)'] = 'TABLA DE CÁLCULO DE CIRCUITOS'
        datos.update({'Hoja N°': i, 'Plano propio': nombre, 'Página (n / m)': f'{i} / {n}', 'Lista de modificaciones': 'NO',
                      'Archivo de salida': f'{nombre}.svg'})
        for h, v in datos.items():
            if h in hdr and v is not None: ws.cell(i + 1, hdr[h]).value = v
    wb.save(xlsx)
    return xlsx


def rotulo(indice, fmt='A1'):
    """Genera el rótulo de la fila `indice` (0 = primera hoja) con el plugin iram-titleblock-plugin.

    Copia de trabajo con sólo esa fila -> titleblocks-generator.py -> svg/. Devuelve (lamina, bloque_mod).
    """
    import openpyxl, shutil
    xlsx = _excel()
    wb = openpyxl.load_workbook(xlsx)
    ws = wb['Láminas']
    hdr = {c.value: c.column for c in ws[1]}
    fila = indice + 2
    nombre = ws.cell(fila, hdr['Plano propio']).value
    work = os.path.join(CACHE, 'rotulo', str(nombre))
    shutil.rmtree(work, ignore_errors=True)
    os.makedirs(work)
    archivo = ws.cell(fila, hdr['Archivo de salida']).value
    if ws.max_row > fila: ws.delete_rows(fila + 1, ws.max_row - fila)
    if fila > 2: ws.delete_rows(2, fila - 2)
    ws.cell(2, hdr['Formato (HOJA)']).value = fmt   # sólo la copia de trabajo: el Excel de src/ no cambia de formato
    wb.save(os.path.join(work, 'input-titleblocks.xlsx'))
    # logos: junto al Excel en src/ (provider-logo-*.svg, client-logo-*.svg); el generador los busca en la carpeta de trabajo
    for f in os.listdir(ROTULO_DIR):
        if f.endswith('.svg') and ('-logo-' in f):
            shutil.copy(os.path.join(ROTULO_DIR, f), work)
    gen = os.path.join(_dir('iram-titleblock-plugin'), 'scripts', 'titleblocks-generator.py')
    subprocess.run([sys.executable, gen], check=True, capture_output=True, cwd=work)   # sin -I: el plugin necesita ezdxf del usuario
    return os.path.join(work, 'svg', archivo), os.path.join(work, 'svg', 'modifications-block-IRAM4508.svg')


def svg_a_dxf(svg_path):
    """DXF con el mismo convertidor del plugin (svg_to_dxf.convertir): mismo nombre, extensión .dxf."""
    p = os.path.join(_dir('iram-titleblock-plugin'), 'scripts')
    if p not in sys.path:
        sys.path.insert(0, p)
    from svg_to_dxf import convertir
    convertir(svg_path)
    return os.path.splitext(svg_path)[0] + '.dxf'
