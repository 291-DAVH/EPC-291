"""Dependencies of the SLD generator, resolved from the epc_symbols and epc_titleblock packages.

- symbol library: epc_symbols (sld-lv-symbol)
- title block and IRAM sheets: epc_titleblock
- calculation formulas: epc_calc (a regular package, imported directly)
"""
import os, sys, tempfile, subprocess
from pathlib import Path

import epc_titleblock as _titleblock
from epc_symbols import library_dir

HERE = Path(__file__).resolve().parent
CACHE = os.path.join(tempfile.gettempdir(), 'epc_sld_work')


def iec_lib():
    """Carpeta con los símbolos SLD BT (_catalog.json + *.svg)."""
    return str(library_dir('sld-lv-symbol'))


_EXAMPLE_TB = HERE.parents[2] / 'examples' / 'sld' / 'titleblock'     # engine/drawing/examples/sld/titleblock


ROTULO_DIR = os.environ.get('SLD_ROTULO_DIR') or (str(_EXAMPLE_TB) if _EXAMPLE_TB.is_dir() else os.path.join(os.getcwd(), 'titleblock'))   # input-titleblocks.xlsx + logos
# Columnas de la plantilla que son datos de ejemplo del plugin: se vacían al crear el Excel del proyecto
_EJEMPLO = ('Empresa:', 'Cliente:', 'Proyectado por', 'Verificado por', 'Validado por', 'Descripción', 'Ubicación')


def _crear_excel(xlsx, datos):
    """Crea input-titleblocks.xlsx desde la plantilla del plugin (sin datos de ejemplo)."""
    import openpyxl
    src = os.path.join(os.path.dirname(_titleblock.__file__), 'templates', 'input-titleblocks.xlsx')
    wb = openpyxl.load_workbook(src)
    ws = wb['Láminas']
    hdr = {c.value: c.column for c in ws[1]}
    for r in range(2, ws.max_row + 1):
        for h, col in hdr.items():
            if h and str(h).startswith(_EJEMPLO):
                ws.cell(r, col).value = None
    # datos que el SLD conoce, sólo para la fila A1 (se escriben una vez, al crear el archivo)
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, hdr['Formato (HOJA)']).value == 'A1':
            for h, v in datos.items():
                if h in hdr and v:
                    ws.cell(r, hdr[h]).value = v
    # 'Archivo de salida' es una fórmula sin valor en caché: el generador la lee con data_only, así que se escribe el valor
    fi, fa, fm = hdr['Formato (HOJA)'], hdr['Archivo de salida'], hdr['Lista de modificaciones']
    for r in range(2, ws.max_row + 1):
        f = ws.cell(r, fi).value
        if f:
            mod = ws.cell(r, fm).value == 'SI'
            ws.cell(r, fa).value = f"sheet-IRAM-{f}{'-mod' if mod else ''}.svg"
    os.makedirs(os.path.dirname(xlsx), exist_ok=True)
    wb.save(xlsx)


def _excel():
    xlsx = os.path.join(ROTULO_DIR, 'input-titleblocks.xlsx')
    if not os.path.isfile(xlsx):
        _crear_excel(xlsx, {})
    return xlsx


def sincronizar(nombres, comunes, fmt='A1'):
    """Deja input-titleblocks.xlsx con UNA fila por hoja (nombres = planos propios de cada hoja).

    Los datos comunes (empresa, cliente, personas...) se toman de la primera fila `fmt` del Excel.
    Cada fila recibe: hoja N°, plano propio, página n / m, denominación y proyecto de esa hoja.
    """
    import openpyxl
    xlsx = _excel()
    wb = openpyxl.load_workbook(xlsx)
    ws = wb['Láminas']
    hdr = {c.value: c.column for c in ws[1]}
    base = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, hdr['Formato (HOJA)']).value == fmt)
    vals = [c.value for c in ws[base]]
    if ws.max_row > 1: ws.delete_rows(2, ws.max_row - 1)
    n = len(nombres)
    for i, nombre in enumerate(nombres, 1):
        for c, v in enumerate(vals, 1): ws.cell(i + 1, c).value = v
        datos = dict(comunes)
        if str(nombre).endswith('_TABLA'):   # hoja de tabla de cálculo
            datos['Denominación (línea 2)'] = 'TABLA DE CÁLCULO DE CIRCUITOS'
        datos.update({'Hoja N°': i, 'Plano propio': nombre, 'Página (n / m)': f'{i} / {n}', 'Lista de modificaciones': 'NO',
                      'Archivo de salida': f'{nombre}.svg'})
        for h, v in datos.items():
            if h in hdr and v is not None: ws.cell(i + 1, hdr[h]).value = v
    wb.save(xlsx)
    return xlsx


def rotulo(indice, fmt='A1'):
    """Genera el rótulo de la fila `indice` (0 = primera hoja) con el plugin iram-titleblock-plugin.

    Copia de trabajo con sólo esa fila -> titleblocks-generator.py -> svg/. Devuelve (lamina, bloque_mod).
    """
    import openpyxl, shutil
    xlsx = _excel()
    wb = openpyxl.load_workbook(xlsx)
    ws = wb['Láminas']
    hdr = {c.value: c.column for c in ws[1]}
    fila = indice + 2
    nombre = ws.cell(fila, hdr['Plano propio']).value
    work = os.path.join(CACHE, 'rotulo', str(nombre))
    shutil.rmtree(work, ignore_errors=True)
    os.makedirs(work)
    archivo = ws.cell(fila, hdr['Archivo de salida']).value
    if ws.max_row > fila: ws.delete_rows(fila + 1, ws.max_row - fila)
    if fila > 2: ws.delete_rows(2, fila - 2)
    ws.cell(2, hdr['Formato (HOJA)']).value = fmt   # sólo la copia de trabajo: el Excel de src/ no cambia de formato
    wb.save(os.path.join(work, 'input-titleblocks.xlsx'))
    # logos: junto al Excel en src/ (provider-logo-*.svg, client-logo-*.svg); el generador los busca en la carpeta de trabajo
    for f in os.listdir(ROTULO_DIR):
        if f.endswith(('.svg', '.png')) and ('-logo-' in f):
            shutil.copy(os.path.join(ROTULO_DIR, f), work)
    subprocess.run([sys.executable, '-m', 'epc_titleblock.titleblocks_generator'], check=True, capture_output=True, cwd=work)   # sin -I: necesita ezdxf del usuario
    return os.path.join(work, 'svg', archivo), os.path.join(work, 'svg', 'modifications-block-IRAM4508.svg')


def svg_a_dxf(svg_path):
    """DXF con el mismo convertidor del paquete (svg_to_dxf.convertir): mismo nombre, extensión .dxf."""
    from epc_titleblock.svg_to_dxf import convertir
    convertir(svg_path)
    return os.path.splitext(svg_path)[0] + '.dxf'
