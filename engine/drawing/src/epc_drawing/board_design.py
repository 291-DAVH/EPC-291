"""Board design: design criteria -> circuit calculation (epc_aea90364) -> single-line diagram (SLD) of a board.

The calculation is the single source of truth: the project's design criteria govern `epc_aea90364.aea_pipeline`, and its result
(In, Ib, Icu, section, accumulated voltage drop, device) is what the SLD generator draws, with the same limits in its notes.

  python -m epc_drawing.board_design --input board.json --out DIR --name TGBT-01 [--criteria criteria.json] [--diagrams sld,mld]
  python -m epc_drawing.board_design --project <project folder> --out DIR --name TGBT-01 [--diagrams sld,mld]   # input and criteria from the project
                                     [--logo provider=FILE] [--logo client=FILE] [--maxl 18]

board.json (the same data the web form saves):
  {"board": {name, voltage_v, system, icc_ka, fp_default, frequency_hz, project, feeder_length_m},
   "circuits": [{tag, description, load, unit, destination, parent_tag, topology, phase, length_m, fp, efficiency, kind, notes, cable_type}],
   "titleblock": {company, company_cuit, client, client_cuit, drawn_by, checked_by, approved_by, description, location, scale, revision, status}}
criteria.json: {criterion id: value}, e.g. {"vd_sectional_pct": 1, "vd_lighting_sockets_pct": 3, "vd_motors_pct": 5, "insulation": "XLPE", ...}

Writes in DIR/out: <NAME>_dimensionado.json (contract of aea_export), SLD_<NAME>*/MLD_<NAME>* sheets, <NAME>_dimensionado_detalle.json (type, limit and drops per circuit),
<NAME>_criterios.json (criteria applied), the SVG/DXF sheets of the SLD generator and a summary (summary.json). The last stdout line is "OK <summary path>".
"""
import argparse
import datetime as dt
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Dict, List, Optional

BOARD_ROWS = [("NOMBRE", "name"), ("TENSION_V", "voltage_v"), ("SISTEMA", "system"), ("ICC_KA", "icc_ka"), ("FP_DEFAULT", "fp_default"),
              ("FRECUENCIA_HZ", "frequency_hz"), ("PROYECTO", "project"), ("LONG_ACOMETIDA_m", "feeder_length_m")]
CIRCUIT_COLS = [("TAG", "tag"), ("DESCRIPCION", "description"), ("CARGA", "load"), ("UNIDAD_CARGA", "unit"), ("DESTINO", "destination"), ("PARENT_TAG", "parent_tag"),
                ("TOPOLOGY", "topology"), ("FASE", "phase"), ("LONGITUD_m", "length_m"), ("FP", "fp"), ("RENDIMIENTO", "efficiency"), ("KIND", "kind"),
                ("NOTAS", "notes"), ("CABLE_TIPO", "cable_type"), ("METODO", "method"), ("AISLACION", "insulation"), ("MATERIAL", "material")]



def catalog_file(discipline: str = "electrical") -> Optional[Path]:
    """The design-criteria catalog: the only place where the base values live (the norms reviewer corrects them there). It is data of the norm
    (`epc_aea90364.norm_data`), so it travels with the engine, in the repository and in a built plugin alike."""
    from epc_aea90364 import norm_data
    return norm_data.criteria_file(discipline)


def catalog_defaults(discipline: str = "electrical") -> dict:
    """Base value of every criterion, read from the catalog."""
    f = catalog_file(discipline)
    if f is None:
        raise FileNotFoundError("design-criteria.json not found: the plugin was built without its catalog")
    return {c["id"]: c["default"] for c in json.loads(f.read_text(encoding="utf-8-sig"))["criteria"]}


def resolve_criteria(project_dir: Optional[Path] = None, discipline: str = "electrical") -> dict:
    """Catalog values with the project's overrides (`workspace/00-setup/design-criteria.json`, written by the web app) applied."""
    crit = catalog_defaults(discipline)
    if project_dir:
        f = Path(project_dir) / "workspace" / "00-setup" / "design-criteria.json"
        if f.is_file():
            ov = json.loads(f.read_text(encoding="utf-8-sig")).get("overrides", {}).get(discipline, {})
            crit.update({k: v["value"] for k, v in ov.items() if isinstance(v, dict) and "value" in v})
    return crit


def _flags(c: dict) -> dict:
    """Marks the pipeline reads that are not Excel columns: a length measured on the plan by the moment method (see aea_pipeline)."""
    return {"MOMENTOS": True} if c.get("moments") else {}


def _board_from_sheet(ws) -> dict:
    values = {row[0]: row[1] for row in ws.iter_rows(values_only=True) if row and row[0]}
    return {key: values[col] for col, key in BOARD_ROWS if values.get(col) is not None}


def _is_circuit_row(r) -> bool:
    return bool(r) and r[0] is not None and not str(r[0]).lower().startswith(("campo", "fila"))


def _circuit_fields(row: dict) -> dict:
    return {key: row[col] for col, key in CIRCUIT_COLS if row.get(col) is not None}


def _circuits_from_sheet(ws) -> List[dict]:
    rows = list(ws.iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r and r[0] == "TAG")
    hdr = [str(h).strip() if h else "" for h in rows[head]]
    named = [dict(zip(hdr, r)) for r in rows[head + 1:] if _is_circuit_row(r)]
    return [_circuit_fields(row) for row in named]


def read_excel(path) -> dict:
    """The workbook of the plugin's example (sheets TABLERO and CIRCUITOS) as input data, the inverse of `write_excel`."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    return {"board": _board_from_sheet(wb["TABLERO"]), "circuits": _circuits_from_sheet(wb["CIRCUITOS"])}


def excel_rows(data: dict):
    """(tablero, circuitos) with the Excel column names that the plugin and aea_pipeline read."""
    tablero = {col: data["board"].get(key) for col, key in BOARD_ROWS}
    circuitos = []
    for c in data["circuits"]:
        row = {col: (None if c.get(key) in (None, "") else c.get(key)) for col, key in CIRCUIT_COLS}
        if c.get("route"):                                           # real route: its total is the length of the cable (single source)
            row["RUTA"] = c["route"]
            row["LONGITUD_m"] = round(sum(float(s["length_m"]) for s in c["route"]), 3)
        circuitos.append({**row, **_flags(c)})
    return tablero, circuitos


def write_excel(path: Path, tablero: dict, circuitos: List[dict]) -> None:
    """The workbook the SLD generator reads: sheets TABLERO (CAMPO/VALOR) and CIRCUITOS."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TABLERO"
    ws.append(["DATOS DEL TABLERO"])
    ws.append(["CAMPO", "VALOR"])
    for col, _ in BOARD_ROWS:
        ws.append([col, tablero.get(col)])
    cs = wb.create_sheet("CIRCUITOS")
    cs.append(["CIRCUITOS"])
    cs.append([c for c, _ in CIRCUIT_COLS])
    for row in circuitos:
        cs.append([row.get(c) for c, _ in CIRCUIT_COLS])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def calculate(data: dict, criteria: Optional[dict] = None) -> dict:
    """aea_pipeline with the design criteria: {"datos": contract JSON, "avisos": [...], "detalle": {...}, "limites": {...}}."""
    from epc_aea90364 import aea_pipeline as P
    tablero, circuitos = excel_rows(data)
    detalle: Dict[str, dict] = {}
    datos, avisos = P.calculate(tablero, circuitos, criterios=criteria, detalle=detalle)
    for c in data.get("circuits", []):                                # the description travels with the calculation, so any view of it reads without the input
        if c.get("tag") in detalle:
            detalle[c["tag"]]["descripcion"] = str(c.get("description", ""))
    limites, _, material, aislacion, sens = P.resolver_criterios(criteria)
    return {"datos": datos, "avisos": avisos, "detalle": detalle,
            "limites": {"seccional": limites["seccional"], "terminal": limites["iluminacion"], "motor": limites["motor"]},
            "conductor": {"material": material, "aislacion": aislacion}, "rcd_ma": sens}


def calc_note(res: dict) -> str:
    lim, cond = res["limites"], res["conductor"]
    mat = {"cu": "cobre", "al": "aluminio"}[cond["material"]]
    return (f"Cálculo con los criterios de diseño del proyecto (AEA 90364): ΔU seccionales {lim['seccional']:g} %, terminales {lim['terminal']:g} %, "
            f"motores {lim['motor']:g} %; {mat}, aislación {cond['aislacion'].upper()}.")


def groups_in(data: dict, criteria: Optional[dict] = None) -> dict:
    """The board data with the differential groups the rules of the standard solutions ask for (epc_aea90364.grouping); the same data when no rule applies."""
    from epc_aea90364 import grouping
    tablero, rows = excel_rows(data)
    groups = grouping.plan(tablero, rows, criteria)
    if not groups:
        return data
    out = json.loads(json.dumps(data))
    parent = {m: g["tag"] for g in groups for m in g["members"]}
    first = {min(g["members"], key=[c["tag"] for c in data["circuits"]].index): g for g in groups}
    circuits = []
    for c in out["circuits"]:
        if c["tag"] in first:
            g = first[c["tag"]]
            circuits.append({"tag": g["tag"], "description": g["description"], "load": None, "unit": "A", "destination": "", "parent_tag": "", "topology": "single", "phase": g["phase"],
                             "length_m": 0, "fp": None, "efficiency": None, "kind": "rccb", "notes": "", "cable_type": "tetra_sep", "group_of": g["members"]})
        circuits.append({**c, "parent_tag": parent[c["tag"]]} if c["tag"] in parent else c)
    out["circuits"] = circuits
    return out


def _groups_of(data: dict) -> dict:
    """The differential groups the standard solutions added to the board: tag -> the circuits behind it."""
    return {c["tag"]: c["group_of"] for c in data["circuits"] if c.get("group_of")}


def _package_index(out: Path, name: str, error) -> None:
    """The index of the delivery package: every sheet of the single-line and the multi-line with its DXF and the layers it uses (epc_dxf)."""
    from epc_dxf import package
    svgs = sorted(out.glob(f"SLD_{name}*.svg")) + sorted(out.glob(f"MLD_{name}*.svg"))
    if svgs and error is None:
        package.describe(svgs, out, name)


def _finish(out: Path, summary: dict) -> dict:
    (out / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    return summary


def _summary(name: str, avisos: list, log: list, res: dict, data: dict, diagrams, memo: Optional[dict], error: Optional[str]) -> dict:
    return {"name": name, "ok": error is None, "avisos": avisos, "log": log, "limites_pct": res["limites"], "conductor": res["conductor"], "circuits": len(res["datos"]), "groups": _groups_of(data),
            "non_compliant": [t for t, d in res["detalle"].items() if not d["cumple"]], "diagrams": list(diagrams), "memo": memo["resumen"] if memo else None, "error": error}


def design(data: dict, out_dir: Path, name: str, criteria: Optional[dict] = None, logos: Optional[Dict[str, Path]] = None, maxl: int = 18,
           diagrams=("sld",)) -> dict:
    out_dir = Path(out_dir)
    out = out_dir / "out"
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    data = groups_in(data, criteria)                          # circuits without a differential are put behind one where the standard solutions ask for it
    res = calculate(data, criteria)
    tablero, circuitos = excel_rows(data)
    xlsx = out_dir / "input_for_sld.xlsx"
    write_excel(xlsx, tablero, circuitos)
    (out_dir / "input_design.json").write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    from epc_titleblock import service as titleblock
    request_json = out_dir / "titleblock-request.json"                  # what the diagrams ask the title block service for: the project's data and logos
    request_json.write_text(json.dumps(titleblock.project_request(data.get("titleblock", {}), logos), ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")

    calc_json = out / f"{name}_dimensionado.json"
    detail_json = out / f"{name}_dimensionado_detalle.json"
    crit_json = out / f"{name}_criterios.json"
    from epc_aea90364 import aea_export as EX
    avisos = list(res["avisos"]) + list(EX.export_calculation(res["datos"], str(calc_json)))
    detail_json.write_text(json.dumps(res["detalle"], ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    crit_json.write_text(json.dumps({"criterios": criteria or {}, "limites_pct": res["limites"], "conductor": res["conductor"], "rcd_ma": res["rcd_ma"]},
                                    ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")

    if "calc" in diagrams:                                    # the calculation alone (sections, protections, drops): no memo, no figures, no diagrams
        return _finish(out, _summary(name, avisos, [], res, data, diagrams, None, None))
    from epc_aea90364 import board_memo as BM                 # memoria de cálculo y síntesis del tablero: mismos números que el unifilar (fuente única)
    memo = BM.build(data["board"], data["circuits"], res["datos"], res["detalle"], res["limites"], res["conductor"], res["rcd_ma"],
                    upstream=data.get("upstream"), validated=bool(data.get("validated", True)))
    (out / f"{name}_memoria.json").write_text(json.dumps(memo, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    (out / f"{name}_MC.md").write_text(BM.memoria_md(memo), encoding="utf-8", newline="\n")
    (out / f"{name}_SP.md").write_text(BM.sintesis_md(memo), encoding="utf-8", newline="\n")
    try:                                                   # the memo and the summary as PDF (what is issued to the client)
        from epc_drawing.export import report_pdf
        report_pdf.render(BM.memoria_doc(memo), out / f"{name}_MC.pdf")
        report_pdf.render(BM.sintesis_doc(memo), out / f"{name}_SP.pdf")
    except Exception as exc:                               # noqa: BLE001 - the calculation stands; the PDF is reported as a warning
        avisos.append(f"no se pudo generar el PDF de la memoria: {exc}")

    from epc_aea90364 import board_figures as BF               # what the drawings show beyond each circuit: bar, feeder, checks and notes (the drawings only draw)
    figures = BF.build(data["board"], res["datos"], res["detalle"], res["limites"], res["rcd_ma"], memo["acometida"])
    figures["note"] += " " + calc_note(res)
    figures_json = out / f"{name}_cifras.json"
    figures_json.write_text(json.dumps(figures, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    env = dict(os.environ, EPC_TITLEBLOCK_JSON=str(request_json), SLD_CALC_JSON=str(calc_json), SLD_CALC_DETAIL=str(detail_json), SLD_FIGURES_JSON=str(figures_json))
    log, error = [], None
    if "sld" in diagrams:
        run = subprocess.run([sys.executable, "-m", "epc_sld.sld_paginate", "default", str(xlsx), str(out), f"SLD_{name}", "default", "default", str(maxl)],
                             env=env, capture_output=True, text=True)
        log += [ln for ln in (run.stdout.splitlines() + run.stderr.splitlines()) if ln.startswith(("OK", "AVISO"))]
        if run.returncode != 0:
            error = (run.stderr or run.stdout).strip()[-600:]
    if "mld" in diagrams and error is None:
        run = subprocess.run([sys.executable, "-m", "epc_mld.mld_generator", "--input", str(out_dir / "input_design.json"), "--calc", str(calc_json),
                              "--detail", str(detail_json), "--out", str(out), "--name", f"MLD_{name}"], env=env, capture_output=True, text=True)
        log += [ln for ln in (run.stdout.splitlines() + run.stderr.splitlines()) if ln.startswith(("OK", "AVISO"))]
        if run.returncode != 0:
            error = (run.stderr or run.stdout).strip()[-600:]
    _package_index(out, name, error)
    return _finish(out, _summary(name, avisos, log, res, data, diagrams, memo, error))


def _load_input(a, ap):
    """The board data and the design criteria the command line asks for: a project, a JSON file or the example workbook."""
    if a.project:
        pkg = Path(a.project) / "workspace" / "02-engineering" / "electrical" / "PQ-01"
        data = json.loads((pkg / "src" / "sld-input.json").read_text(encoding="utf-8"))
        criteria = resolve_criteria(Path(a.project))
    elif a.xlsx:
        data, criteria = read_excel(a.xlsx), None
    elif a.input:
        data = json.loads(Path(a.input).read_text(encoding="utf-8"))
        criteria = json.loads(Path(a.criteria).read_text(encoding="utf-8")) if a.criteria else None
    else:
        ap.error("give --project, --input or --xlsx")
    return data, criteria


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Criteria-driven calculation and single-line diagram of a board")
    ap.add_argument("--project", help="project folder: input from its saved SLD data and criteria from its design criteria (instead of --input/--criteria)")
    ap.add_argument("--input")
    ap.add_argument("--xlsx", help="the plugin's example workbook as input (instead of --input)")
    ap.add_argument("--out", required=True)
    ap.add_argument("--name", required=True)
    ap.add_argument("--criteria")
    ap.add_argument("--logo", action="append", default=[], help="provider=FILE or client=FILE")
    ap.add_argument("--maxl", type=int, default=18)
    ap.add_argument("--diagrams", default="sld", help="comma list: sld,mld (or none: calculation, memo and summary only; or calc: the calculation alone, no memo)")
    a = ap.parse_args(argv)
    data, criteria = _load_input(a, ap)
    logos = {k: Path(v) for k, v in (x.split("=", 1) for x in a.logo)}
    s = design(data, Path(a.out), a.name, criteria, logos, a.maxl, tuple(x for x in a.diagrams.split(",") if x))
    for ln in s["log"] + s["avisos"]:
        print(ln, file=sys.stderr)
    if not s["ok"]:
        print("FAIL", s["error"], file=sys.stderr)
        return 1
    print("OK", Path(a.out) / "out" / "summary.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
